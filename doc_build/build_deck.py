# -*- coding: utf-8 -*-
"""Monta a documentação técnica do firmware BLE Mesh sobre o template oficial
Microchip (assets/Microchip_template.pptx da skill microchip-pptx-template).

Parte do template já copiado em ble_mesh_doc.pptx: mantém o slide de título
(master 1, layout '1_Title Slide'), descarta os slides de exemplo e cria o
conteúdo sobre o layout 'Title Only' do master 4 — o mesmo master do único
slide de conteúdo bom do template, cujo rodapé e logo vêm de shapes do layout
e portanto são herdados automaticamente.

Paleta e fonte saem do próprio tema do template (Calibri, navy 0E3689,
ciano 1D9CE4, laranja FD7F20).
"""
import copy

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

DECK = 'ble_mesh_doc.pptx'

# ---------------------------------------------------------------- paleta
NAVY = RGBColor(0x0E, 0x36, 0x89)
CYAN = RGBColor(0x1D, 0x9C, 0xE4)
ORANGE = RGBColor(0xFD, 0x7F, 0x20)
GREEN = RGBColor(0x5E, 0xBF, 0x33)
INK = RGBColor(0x0A, 0x0B, 0x0F)
MUTED = RGBColor(0x5B, 0x66, 0x7A)
CARD = RGBColor(0xF1, 0xF4, 0xFA)
CARD_ALT = RGBColor(0xFF, 0xF4, 0xE8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RULE = RGBColor(0xD6, 0xDE, 0xEC)

BODY_FONT = 'Calibri'
MONO_FONT = 'Courier New'

TOP = 1.42
BOT = 6.88
LEFT = 0.55
RIGHT = 12.78
WIDTH = RIGHT - LEFT


# ---------------------------------------------------------------- helpers
def set_text(tf, blocks, default_size=14, default_color=INK, line=0.95,
             space_after=4, font=BODY_FONT):
    # Obrigatorio: os placeholders do slide de titulo do template ja vem com
    # texto ('Rodrigo Britto', 'April 2026'). Sem limpar, add_run() concatena
    # nos runs existentes e o slide sai com os dois valores sobrepostos.
    tf.clear()
    tf.word_wrap = True
    first = True
    for block in blocks:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        if isinstance(block, dict):
            block = [block]
        elif isinstance(block, str):
            block = [{'t': block}]
        opts = {}
        if block and isinstance(block[0], dict) and '_' in block[0]:
            opts = block[0].pop('_')
            if not block[0].get('t'):
                block = block[1:]
        p.alignment = opts.get('align', PP_ALIGN.LEFT)
        p.space_after = Pt(opts.get('after', space_after))
        p.space_before = Pt(opts.get('before', 0))
        p.line_spacing = opts.get('line', line)
        for run in block:
            r = p.add_run()
            r.text = run['t']
            f = r.font
            f.name = run.get('font', opts.get('font', font))
            f.size = Pt(run.get('size', opts.get('size', default_size)))
            f.bold = run.get('b', opts.get('b', False))
            f.italic = run.get('i', False)
            f.color.rgb = run.get('c', opts.get('c', default_color))
    return tf


def textbox(slide, x, y, w, h, blocks, anchor=MSO_ANCHOR.TOP, **kw):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    set_text(tf, blocks, **kw)
    return box


def card(slide, x, y, w, h, fill=CARD, line=None, radius=0.10):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = radius
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    sh.shadow.inherit = False
    sh.text_frame.text = ''
    return sh


def badge(slide, x, y, d, label, fill=NAVY, color=WHITE, size=13):
    """Círculo pequeno com rótulo — o motivo visual repetido no deck."""
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y),
                                Inches(d), Inches(d))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.name = BODY_FONT
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = color
    return sh


def chip(slide, x, y, w, h, label, fill, color=WHITE, size=11, bold=True,
         font=BODY_FONT, radius=0.5):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = radius
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return sh


def connect(slide, x1, y1, x2, y2, color=RULE, width=1.5):
    from pptx.enum.shapes import MSO_CONNECTOR
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1),
                                   Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    return c


def table(slide, x, y, w, h, rows, col_w, header_fill=NAVY, size=11.5,
          header_size=11.5, first_col_bold=False, mono_cols=()):
    n_rows, n_cols = len(rows), len(rows[0])
    gf = slide.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y),
                                Inches(w), Inches(h))
    tbl = gf.table
    tbl.first_row = True
    tbl.horz_banding = True
    for i, cw in enumerate(col_w):
        tbl.columns[i].width = Inches(cw)
    for ri, row in enumerate(rows):
        tbl.rows[ri].height = Inches(h / n_rows)
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.margin_left = Inches(0.09)
            cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if ri == 0:
                cell.fill.fore_color.rgb = header_fill
            else:
                cell.fill.fore_color.rgb = WHITE if ri % 2 else CARD
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            r = p.add_run()
            r.text = val
            r.font.name = MONO_FONT if (ri > 0 and ci in mono_cols) else BODY_FONT
            r.font.size = Pt(header_size if ri == 0 else size)
            r.font.bold = (ri == 0) or (first_col_bold and ci == 0)
            r.font.color.rgb = WHITE if ri == 0 else INK
    return tbl


# ------------------------------------------------------- montagem do deck
prs = Presentation(DECK)
masters = list(prs.slide_masters)
LAY = [l for l in masters[4].slide_layouts if l.name == 'Title Only'][0]

sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst)[1:]:
    rId = sldId.get(
        '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
    prs.part.drop_rel(rId)
    sldIdLst.remove(sldId)

_num_ph = None
for ph in LAY.placeholders:
    if ph.placeholder_format.idx == 4:
        _num_ph = ph._element


def new_slide(title, kicker=None):
    s = prs.slides.add_slide(LAY)
    if _num_ph is not None:
        s.shapes._spTree.append(copy.deepcopy(_num_ph))
    t = s.shapes.title
    t.left, t.top, t.width, t.height = (Inches(LEFT), Inches(0.34),
                                        Inches(WIDTH), Inches(0.78))
    tf = t.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    set_text(tf, [[{'t': title, 'size': 30, 'b': True, 'c': NAVY}]])
    if kicker:
        textbox(s, LEFT, 1.10, WIDTH, 0.26,
                [[{'t': kicker, 'size': 12, 'c': MUTED, 'i': True}]])
    return s


# ================================================================ slide 1
s1 = prs.slides[0]
ph = {p.name: p for p in s1.placeholders}
set_text(ph['Title 3'].text_frame,
         [[{'t': 'Firmware BLE Mesh para Dimmers', 'size': 34, 'b': True,
            'c': NAVY}],
          [{'_': {'before': 6},
            't': 'Protocolo, topologia e tempos de formação de rede',
            'size': 19, 'c': CYAN}]])
set_text(ph['Text Placeholder 4'].text_frame,
         [[{'_': {'align': PP_ALIGN.RIGHT}, 't': 'Rodrigo Britto', 'size': 18,
            'b': True, 'c': NAVY}]])
set_text(ph['Text Placeholder 1'].text_frame,
         [[{'_': {'align': PP_ALIGN.RIGHT}, 't': 'Agosto de 2026', 'size': 18,
            'c': CYAN}]])

# ================================================================ slide 2
s = new_slide('Escopo e plataforma')
textbox(s, LEFT, TOP, 12.0, 0.75, [[
    {'t': 'Malha BLE proprietária sobre GATT/TRSP', 'b': True, 'size': 15},
    {'t': '  —  cada nó é um dimmer RGB que também roteia tráfego para os '
          'vizinhos. Não é BLE Mesh SIG: o encaminhamento é feito por '
          'inundação controlada sobre conexões GATT persistentes, o que '
          'mantém a latência baixa e dispensa provisionamento de chaves.',
     'size': 15, 'c': INK}]], line=1.05)

stats = [('100', 'nós endereçáveis', 'NODE_ID 1..100'),
         ('6', 'enlaces por nó', '1 central + 5 periféricos'),
         ('20 B', 'pacote mesh', 'cabeçalho de 5 B'),
         ('4', 'níveis de árvore', 'diâmetro de 8 saltos')]
cw, gap = 2.94, 0.16
for i, (big, lab, sub) in enumerate(stats):
    x = LEFT + i * (cw + gap)
    card(s, x, 2.32, cw, 1.42)
    textbox(s, x + 0.22, 2.46, cw - 0.44, 0.55,
            [[{'t': big, 'size': 30, 'b': True, 'c': NAVY}]])
    textbox(s, x + 0.22, 3.03, cw - 0.44, 0.28,
            [[{'t': lab, 'size': 13, 'b': True, 'c': INK}]])
    textbox(s, x + 0.22, 3.33, cw - 0.44, 0.28,
            [[{'t': sub, 'size': 11, 'c': MUTED}]])

cw2 = (WIDTH - 0.24) / 2
card(s, LEFT, 4.02, cw2, 2.48)
# rotulos de duas letras precisam de circulo maior e corpo menor: 'HW' a 13pt
# quebrava em duas linhas dentro de um circulo de 0.34".
badge(s, LEFT + 0.24, 4.22, 0.40, 'HW', size=11)
textbox(s, LEFT + 0.74, 4.28, cw2 - 0.98, 0.3,
        [[{'t': 'Plataformas', 'size': 15, 'b': True, 'c': NAVY}]])
textbox(s, LEFT + 0.24, 4.76, cw2 - 0.48, 1.6, [
    [{'t': 'WBZ451', 'b': True}, {'t': '  (PIC32CX-BZ2) — projeto '
                                       'BLE_Mesh, build CMake/Ninja'}],
    [{'t': 'PIC32CX-BZ6', 'b': True}, {'t': '  — projeto BLE_MESH_BZ6, '
                                            'mesma base de código mesh'}],
    [{'t': 'Harmony v3 + FreeRTOS', 'b': True},
     {'t': '  — stack BLE da Microchip'}],
    [{'t': 'Núcleo mesh idêntico', 'b': True},
     {'t': ' nos dois alvos; diverge só a camada de plataforma'}],
], default_size=13, line=1.0, space_after=7)

card(s, LEFT + cw2 + 0.24, 4.02, cw2, 2.48)
badge(s, LEFT + cw2 + 0.48, 4.22, 0.40, 'FN', fill=CYAN, size=11)
textbox(s, LEFT + cw2 + 0.98, 4.28, cw2 - 0.98, 0.3,
        [[{'t': 'Papéis lógicos', 'size': 15, 'b': True, 'c': NAVY}]])
textbox(s, LEFT + cw2 + 0.48, 4.76, cw2 - 0.48, 1.6, [
    [{'t': 'Nó dimmer', 'b': True},
     {'t': '  — saída RGB, roteador e repetidor da malha'}],
    [{'t': 'Raiz (root)', 'b': True},
     {'t': '  — menor node id alcançável; agrega a topologia'}],
    [{'t': 'App / GUI', 'b': True},
     {'t': '  — celular conectado a um nó qualquer, com slot reservado'}],
    [{'t': 'Cluster', 'b': True},
     {'t': '  — grupo de 10 nós, endereçável em bloco'}],
], default_size=13, line=1.0, space_after=7)

# ================================================================ slide 3
s = new_slide('Arquitetura de software',
              'Camadas do firmware e cadência das tarefas periódicas')
layers = [
    ('app.c — APP_Tasks', 'máquina de estados, fila de eventos, cadências',
     NAVY, WHITE),
    ('mesh_routing.c', 'roteamento, dedup, ACK/retry, heartbeat, topologia',
     CYAN, WHITE),
    ('mesh_conn_mgr.c', 'tabela de conexões, papéis, slots, varredura',
     RGBColor(0x4F, 0x6C, 0xB0), WHITE),
    ('app_ble/* (Harmony)', 'GAP, GATT, TRSPS/TRSPC, callbacks',
     RGBColor(0xC9, 0xD5, 0xEC), INK),
    ('BLE stack + FreeRTOS', 'controlador, escalonamento de rádio',
     RGBColor(0xE6, 0xEC, 0xF7), INK),
]
lx, lw = LEFT, 6.55
for i, (name, desc, fill, fg) in enumerate(layers):
    y = TOP + 0.28 + i * 1.02
    card(s, lx, y, lw, 0.86, fill=fill)
    textbox(s, lx + 0.26, y + 0.13, lw - 0.5, 0.30,
            [[{'t': name, 'size': 14, 'b': True, 'c': fg}]])
    textbox(s, lx + 0.26, y + 0.45, lw - 0.5, 0.28,
            [[{'t': desc, 'size': 11.5,
               'c': WHITE if fg == WHITE else MUTED}]])

rx = LEFT + lw + 0.35
rw = RIGHT - rx
card(s, rx, TOP + 0.28, rw, 2.42, fill=CARD_ALT)
badge(s, rx + 0.26, TOP + 0.50, 0.34, '⏱', fill=ORANGE, size=15)
textbox(s, rx + 0.70, TOP + 0.54, rw - 0.96, 0.3,
        [[{'t': 'Duas cadências', 'size': 15, 'b': True, 'c': NAVY}]])
textbox(s, rx + 0.26, TOP + 1.02, rw - 0.52, 1.55, [
    [{'t': '250 ms — MESH_Maintenance()', 'b': True, 'c': ORANGE}],
    [{'t': 'Fila de envio diferido, retransmissão confiável, JOIN, '
           'heartbeat e relatório de topologia. É o piso da latência por '
           'salto da malha.', 'size': 12}],
    [{'_': {'before': 5}, 't': '1 s — descoberta e supervisão',
      'b': True, 'c': NAVY}],
    [{'t': 'Rescan, tentativa de conexão e varredura de enlaces mortos. '
           'Emite comandos HCI: acelerar não faz convergir mais rápido.',
      'size': 12}],
], default_size=13, line=1.0, space_after=2)

card(s, rx, TOP + 2.95, rw, 2.35)
badge(s, rx + 0.26, TOP + 3.15, 0.40, 'RX', fill=NAVY, size=11)
textbox(s, rx + 0.76, TOP + 3.21, rw - 1.02, 0.3,
        [[{'t': 'Caminho de recepção', 'size': 15, 'b': True, 'c': NAVY}]])
textbox(s, rx + 0.26, TOP + 3.69, rw - 0.52, 1.5, [
    [{'t': 'Todo pacote enfileirado no TRSPS é drenado no mesmo evento. O '
           'crédito CBFC só retorna em BLE_TRSPS_GetData(): um pacote '
           'deixado na fila custa um crédito em definitivo, e depois de 16 '
           'o enlace fica mudo em uma das direções.', 'size': 12.5}],
], line=1.05)

# ================================================================ slide 4
s = new_slide('Topologia: regras de formação',
              'Árvore com uma única aresta de subida por nó')
tx, ty = LEFT + 0.1, TOP + 0.55
nodes_l1 = [(tx + 2.55, ty)]
nodes_l2 = [(tx + 0.95, ty + 1.25), (tx + 2.15, ty + 1.25),
            (tx + 3.35, ty + 1.25)]
nodes_l3 = [(tx + 0.35, ty + 2.50), (tx + 1.55, ty + 2.50),
            (tx + 2.75, ty + 2.50), (tx + 3.95, ty + 2.50)]
D = 0.62
for (px, py) in nodes_l1:
    for (cx, cy) in nodes_l2:
        connect(s, px + D / 2, py + D, cx + D / 2, cy, RULE, 1.5)
for (cx, cy) in nodes_l3:
    par = nodes_l2[0] if cx < tx + 1.4 else (nodes_l2[1] if cx < tx + 2.6
                                             else nodes_l2[2])
    connect(s, par[0] + D / 2, par[1] + D, cx + D / 2, cy, RULE, 1.5)
for (px, py) in nodes_l1:
    badge(s, px, py, D, '01', fill=NAVY, size=13)
for i, (px, py) in enumerate(nodes_l2):
    badge(s, px, py, D, '%02d' % (i + 2), fill=CYAN, size=13)
for i, (px, py) in enumerate(nodes_l3):
    badge(s, px, py, D, '%02d' % (i + 5), fill=RGBColor(0x8F, 0xA9, 0xD6),
          size=13)
phx, phy = tx + 4.55, ty + 1.25
connect(s, nodes_l2[2][0] + D, nodes_l2[2][1] + D / 2, phx, phy + 0.20,
        ORANGE, 1.5)
chip(s, phx, phy, 0.92, 0.40, 'APP', ORANGE, size=12)
textbox(s, tx, ty + 3.28, 5.6, 0.9, [
    [{'t': 'depth 0 (raiz)', 'c': NAVY, 'b': True, 'size': 11},
     {'t': '   •   ', 'c': RULE, 'size': 11},
     {'t': 'depth 1', 'c': CYAN, 'b': True, 'size': 11},
     {'t': '   •   ', 'c': RULE, 'size': 11},
     {'t': 'depth 2', 'c': RGBColor(0x6F, 0x88, 0xB6), 'b': True, 'size': 11},
     {'t': '   •   ', 'c': RULE, 'size': 11},
     {'t': 'slot do app', 'c': ORANGE, 'b': True, 'size': 11}],
    [{'t': 'A raiz é o menor node id alcançável. root e depth propagam por '
           'JOIN_ACCEPT e heartbeat.', 'size': 11, 'c': MUTED}],
], space_after=3)

rx = LEFT + 5.95
rw = RIGHT - rx
rules = [
    ('1', 'Uma aresta de subida', 'MESH_MAX_CENTRAL = 1. Cada nó abre no '
     'máximo uma conexão como central, o que torna a direção determinística '
     'e evita tempestades de conexão recíproca.'),
    ('2', 'Só para node id menor', 'Um nó só sobe para um id inferior. A '
     'relação pai é estritamente decrescente, o que impede ciclo e a '
     'contagem ao infinito de depth.'),
    ('3', 'Quatro filhos + o app', 'MESH_MAX_PERIPHERAL = 5, dos quais 4 '
     'são filhos de malha. O quinto slot fica reservado ao celular e não '
     'pode ser tomado por um nó.'),
    ('4', 'Subida excepcional', 'Um nó isolado pode abrir aresta para id '
     'maior. Sem isso os slots dos ids mais altos ficam inalcançáveis e a '
     'rede se parte.'),
]
yy = TOP + 0.30
for num, head, body in rules:
    card(s, rx, yy, rw, 1.22)
    badge(s, rx + 0.22, yy + 0.20, 0.32, num, fill=NAVY, size=13)
    textbox(s, rx + 0.64, yy + 0.19, rw - 0.88, 0.28,
            [[{'t': head, 'size': 14, 'b': True, 'c': NAVY}]])
    textbox(s, rx + 0.64, yy + 0.53, rw - 0.88, 0.62,
            [[{'t': body, 'size': 11.5, 'c': INK}]], line=1.0)
    yy += 1.35

# ================================================================ slide 5
s = new_slide('Topologia: capacidade e profundidade',
              'Quantos nós cabem por nível com fator de ramificação 4')
rows = [['Profundidade', 'Nós no nível', 'Total acumulado', 'Observação'],
        ['0 (raiz)', '1', '1', 'menor node id alcançável'],
        ['1', '4', '5', 'filhos diretos da raiz'],
        ['2', '16', '21', 'cobre uma bancada típica'],
        ['3', '64', '85', 'ainda abaixo de 100 nós'],
        ['4', '256', '341', 'profundidade usada por 100 nós']]
table(s, LEFT, TOP + 0.22, 7.30, 2.62, rows, [1.55, 1.55, 1.85, 2.35],
      first_col_bold=True, size=12)

textbox(s, LEFT, TOP + 3.10, 7.30, 1.9, [
    [{'t': 'É o diâmetro, não a profundidade, que dimensiona o TTL.',
      'b': True, 'size': 14, 'c': NAVY}],
    [{'t': 'mesh_ForwardAll inunda por todos os enlaces, inclusive as '
           'arestas cruzadas, então o alcance necessário é o diâmetro do '
           'grafo de conexões. Com 100 nós em profundidade 4, folha a folha '
           'são 8 saltos. MESH_MAX_TTL = 32 dá margem de 4×.', 'size': 12.5}],
    [{'_': {'before': 6}, 't': 'Exceção: implantação fisicamente linear. ',
      'b': True, 'size': 12.5, 'c': ORANGE},
     {'t': 'Num corredor de 100 luminárias o diâmetro é o próprio '
           'comprimento da cadeia, e o TTL precisa excedê-lo.',
      'size': 12.5}],
], line=1.05, space_after=3)

rx = LEFT + 7.65
rw = RIGHT - rx
card(s, rx, TOP + 0.22, rw, 4.78, fill=CARD_ALT)
badge(s, rx + 0.26, TOP + 0.46, 0.34, '!', fill=ORANGE, size=15)
textbox(s, rx + 0.70, TOP + 0.50, rw - 0.96, 0.3,
        [[{'t': 'Limites estruturais', 'size': 15, 'b': True, 'c': NAVY}]])
textbox(s, rx + 0.26, TOP + 1.00, rw - 0.52, 3.85, [
    [{'t': 'Sem redundância de pai', 'b': True, 'c': NAVY}],
    [{'t': 'Com uma única aresta de subida, a perda de um nó deixa órfã '
           'toda a sua subárvore de uma vez. É escolha de topologia, não '
           'de dimensionamento.', 'size': 12}],
    [{'_': {'before': 8}, 't': 'seq_num de 8 bits', 'b': True, 'c': NAVY}],
    [{'t': 'Um octeto por origem no cabeçalho de fio. Com 100 nós e tráfego '
           'pesado o espaço de sequência fica apertado; alargá-lo quebra a '
           'compatibilidade com a GUI.', 'size': 12}],
    [{'_': {'before': 8}, 't': 'Fila de eventos do app', 'b': True,
      'c': NAVY}],
    [{'t': 'APP_Msg_T ocupa 260 B, e a fila de 64 consome 16,6 KB dos 40 KB '
           'de heap do FreeRTOS — 41% do heap num único objeto.',
      'size': 12}],
], line=1.02, space_after=2)

# ================================================================ slide 6
s = new_slide('Protocolo: formato do pacote',
              'Cabeçalho fixo de 5 bytes seguido de payload opcional')
fields = [('dst_id', '1 B', 'destino'),
          ('src_id', '1 B', 'origem'),
          ('seq_num', '1 B', 'sequência'),
          ('ttl', '1 B', 'saltos restantes'),
          ('cmd', '1 B', 'comando')]
fw, fgap = 1.42, 0.10
fx = LEFT
for name, size_, desc in fields:
    card(s, fx, TOP + 0.30, fw, 1.16, fill=NAVY)
    textbox(s, fx + 0.06, TOP + 0.46, fw - 0.12, 0.3,
            [[{'_': {'align': PP_ALIGN.CENTER}, 't': name, 'size': 13,
               'b': True, 'c': WHITE, 'font': MONO_FONT}]])
    textbox(s, fx + 0.06, TOP + 0.78, fw - 0.12, 0.26,
            [[{'_': {'align': PP_ALIGN.CENTER}, 't': size_, 'size': 11,
               'c': RGBColor(0xBB, 0xCC, 0xEE)}]])
    textbox(s, fx + 0.04, TOP + 1.04, fw - 0.08, 0.26,
            [[{'_': {'align': PP_ALIGN.CENTER}, 't': desc, 'size': 10.5,
               'c': RGBColor(0xBB, 0xCC, 0xEE)}]])
    fx += fw + fgap
pw = RIGHT - fx
card(s, fx, TOP + 0.30, pw, 1.16, fill=CARD, line=RULE)
textbox(s, fx + 0.10, TOP + 0.58, pw - 0.20, 0.3,
        [[{'_': {'align': PP_ALIGN.CENTER}, 't': 'payload', 'size': 13,
           'b': True, 'c': NAVY, 'font': MONO_FONT}]])
textbox(s, fx + 0.10, TOP + 0.90, pw - 0.20, 0.3,
        [[{'_': {'align': PP_ALIGN.CENTER}, 't': '0 a 15 B', 'size': 11,
           'c': MUTED}]])

textbox(s, LEFT, TOP + 1.62, WIDTH, 0.3, [
    [{'t': 'MESH_MAX_PACKET_SIZE = 20 B', 'b': True, 'size': 12.5,
      'c': ORANGE, 'font': MONO_FONT},
     {'t': '   cabe inteiro em uma notificação GATT, sem fragmentação.',
      'size': 12.5}]])

cw3 = (WIDTH - 0.32) / 3
info = [
    ('seq_num', 'Anti-replay por origem', 'Cache indexado por node id (101 '
     'entradas), com um bitmap de 64 posições por origem. Entradas expiram '
     'em 10 s, e a aritmética é em int8_t para tolerar o wrap do contador.'),
    ('ttl', 'Raio da inundação', 'Decrementado a cada salto; em zero o '
     'pacote morre. Não é o TTL que impede repetição — isso é papel do '
     'cache de duplicatas. REPAIR_REQUEST sai com TTL 2, de propósito.'),
    ('cmd', 'Semântica do payload', 'Determina o formato dos bytes '
     'seguintes e se o comando exige ACK. SET_DIMMER, GET_STATUS, IDENTIFY, '
     'PROVISION e UNPROVISION são confirmados.'),
]
for i, (mono, head, body) in enumerate(info):
    x = LEFT + i * (cw3 + 0.16)
    card(s, x, TOP + 2.10, cw3, 2.20)
    chip(s, x + 0.22, TOP + 2.30, 1.10, 0.30, mono, NAVY, size=11,
         font=MONO_FONT)
    textbox(s, x + 0.22, TOP + 2.70, cw3 - 0.44, 0.30,
            [[{'t': head, 'size': 14, 'b': True, 'c': NAVY}]])
    textbox(s, x + 0.22, TOP + 3.06, cw3 - 0.44, 1.15,
            [[{'t': body, 'size': 12}]], line=1.03)

# ================================================================ slide 7
s = new_slide('Protocolo: comandos',
              'Quinze códigos em dois grupos funcionais')
rows_app = [['Cód.', 'Comando', 'Função'],
            ['0x01', 'SET_DIMMER', 'define R, G, B do nó destino'],
            ['0x02', 'GET_STATUS', 'consulta cor e número de enlaces'],
            ['0x03', 'STATUS_RESP', 'resposta de estado (4 B)'],
            ['0x04', 'IDENTIFY', 'pisca o LED para localizar o nó'],
            ['0x05', 'PROVISION', 'grava node id e reinicia'],
            ['0x06', 'UNPROVISION', 'apaga node id e reinicia'],
            ['0x07', 'ACK', 'confirma origem e seq_num']]
rows_mesh = [['Cód.', 'Comando', 'Função'],
             ['0x08', 'HELLO', 'anúncio de identidade no enlace'],
             ['0x09', 'JOIN_REQUEST', 'pede admissão como filho'],
             ['0x0A', 'JOIN_ACCEPT', 'admite e informa root/depth'],
             ['0x0B', 'JOIN_REDIRECT', 'recusa e sugere outro pai'],
             ['0x0C', 'HEARTBEAT', 'presença + root/depth + token'],
             ['0x0F', 'HEARTBEAT_ACK', 'devolve o token do enlace'],
             ['0x0D', 'TOPOLOGY_REPORT', 'agregado da subárvore ao pai'],
             ['0x0E', 'REPAIR_REQUEST', 'aviso de perda de pai (TTL 2)']]
tw = (WIDTH - 0.34) / 2
textbox(s, LEFT, TOP + 0.16, tw, 0.28,
        [[{'t': 'Aplicação e controle', 'size': 14, 'b': True, 'c': NAVY}]])
table(s, LEFT, TOP + 0.52, tw, 3.05, rows_app, [0.85, 2.05, 3.05],
      size=11.5, mono_cols=(0, 1))
textbox(s, LEFT + tw + 0.34, TOP + 0.16, tw, 0.28,
        [[{'t': 'Formação e manutenção da malha', 'size': 14, 'b': True,
           'c': CYAN}]])
table(s, LEFT + tw + 0.34, TOP + 0.52, tw, 3.45, rows_mesh,
      [0.85, 2.30, 2.80], header_fill=CYAN, size=11.5, mono_cols=(0, 1))

textbox(s, LEFT, TOP + 4.20, WIDTH, 0.9, [
    [{'t': 'Confirmação seletiva.', 'b': True, 'size': 12.5, 'c': NAVY},
     {'t': '  Só o unicast dos comandos de aplicação entra na fila '
           'confiável (MESH_TX_QUEUE_SIZE = 24, ACK em 2 s, até 3 '
           'retransmissões com backoff exponencial). Broadcast e cluster '
           'são best effort de propósito: o remetente não tem a lista de '
           'membros, e o GATT já garante entrega em cada enlace físico.',
      'size': 12.5}],
], line=1.05)

# ================================================================ slide 8
s = new_slide('Protocolo: endereçamento',
              'Quatro faixas no campo dst_id de um byte')
addrs = [
    ('0x01 – 0x64', 'Unicast', 'Node id de 1 a 100. Entregue ao nó exato; '
     'gera ACK quando o comando exige.', NAVY),
    ('0x00', 'App / GUI', 'Reservado ao celular. O nó que recebe reescreve '
     'a origem com o próprio id antes de injetar na malha.', ORANGE),
    ('0xC0 – 0xC9', 'Cluster', 'Dez grupos de dez nós: GET_CLUSTER(id) = '
     '(id−1)/10. Cobre exatamente os 100 nós endereçáveis.', CYAN),
    ('0xFF', 'Broadcast', 'Todos os nós. Caminho recomendado para cenas: '
     'não ocupa a fila de ACK.', GREEN),
]
cw4 = (WIDTH - 0.48) / 4
for i, (rng, head, body, col) in enumerate(addrs):
    x = LEFT + i * (cw4 + 0.16)
    card(s, x, TOP + 0.28, cw4, 2.40)
    chip(s, x + 0.22, TOP + 0.52, cw4 - 0.44, 0.38, rng, col, size=12.5,
         font=MONO_FONT)
    textbox(s, x + 0.22, TOP + 1.04, cw4 - 0.44, 0.32,
            [[{'t': head, 'size': 16, 'b': True, 'c': NAVY}]])
    textbox(s, x + 0.22, TOP + 1.46, cw4 - 0.44, 0.85,
            [[{'t': body, 'size': 12}]], line=1.03)

card(s, LEFT, TOP + 2.95, WIDTH, 1.55, fill=CARD_ALT)
badge(s, LEFT + 0.26, TOP + 3.18, 0.34, '→', fill=ORANGE, size=15)
textbox(s, LEFT + 0.70, TOP + 3.21, WIDTH - 1.0, 0.3,
        [[{'t': 'Caminho de um comando vindo do app', 'size': 15, 'b': True,
           'c': NAVY}]])
textbox(s, LEFT + 0.26, TOP + 3.66, WIDTH - 0.52, 0.75, [
    [{'t': 'O nó conectado ao celular reescreve src_id com o próprio id, '
           'atribui novo seq_num e TTL cheio, e registra o pacote no próprio '
           'cache de duplicatas antes de inundar. Sem essa reescrita o '
           'pacote voltaria ao remetente e seria reexecutado; com ela, uma '
           'rajada de slider não se multiplica pela rede.', 'size': 12.5}],
], line=1.05)

# ================================================================ slide 9
s = new_slide('Entrada na rede (JOIN)',
              'Do anúncio até o enlace utilizável para encaminhamento')
steps = [
    ('1', 'Anúncio', 'Cada nó anuncia DIMMER_xx mais dados de fabricante '
     '0xFEDA com flags, slots livres, root e depth.'),
    ('2', 'Descoberta', 'Janela de scan de 6 s a 20% de duty. Candidatos '
     'ordenados por slots livres, root menor, depth menor, falhas e RSSI.'),
    ('3', 'Conexão', 'BLE_GAP_CreateConnection para o melhor candidato de '
     'id inferior. Uma guarda de 15 s cancela a tentativa presa.'),
    ('4', 'GATT + CBFC', 'Descoberta de serviços e handshake TRSPC. Só a '
     'primeira escrita aceita prova que a direção está aberta.'),
    ('5', 'JOIN_REQUEST', 'Enviado com id, root e depth. Retentado com '
     'backoff exponencial até o enlace ficar topologyReady.'),
    ('6', 'JOIN_ACCEPT', 'O pai admite se há slot, devolve root e depth e '
     'atualiza o anúncio. Sem slot, responde JOIN_REDIRECT e desconecta.'),
]
sw = (WIDTH - 0.50) / 3
for i, (num, head, body) in enumerate(steps):
    col, row = i % 3, i // 3
    x = LEFT + col * (sw + 0.25)
    y = TOP + 0.34 + row * 2.30
    card(s, x, y, sw, 2.02)
    badge(s, x + 0.24, y + 0.24, 0.40, num, fill=NAVY if row == 0 else CYAN,
          size=15)
    textbox(s, x + 0.74, y + 0.30, sw - 1.0, 0.32,
            [[{'t': head, 'size': 15, 'b': True, 'c': NAVY}]])
    textbox(s, x + 0.24, y + 0.82, sw - 0.48, 1.05,
            [[{'t': body, 'size': 12.5}]], line=1.05)
    if col < 2:
        textbox(s, x + sw + 0.02, y + 0.86, 0.22, 0.3,
                [[{'_': {'align': PP_ALIGN.CENTER}, 't': '›', 'size': 22,
                   'b': True, 'c': RULE}]])

# ================================================================ slide 10
s = new_slide('Confiabilidade', 'Três mecanismos independentes')
mech = [
    ('Anti-duplicata', NAVY, [
        ('Cache por origem', 'Indexado por node id (101 entradas), não um '
         'FIFO compartilhado: um vizinho falante não despeja o histórico '
         'dos outros.'),
        ('Janela de 64', 'Bitmap uint64_t com aritmética de sequência em '
         'int8_t; 64 é o teto dos dois. Entradas expiram em 10 s.'),
    ]),
    ('Entrega confiável', CYAN, [
        ('Fila de 24 comandos', 'Só unicast de aplicação. Um RGB pode '
         'despejar tráfego administrativo, nunca outro RGB.'),
        ('ACK em 2 s', 'Até 3 retransmissões com backoff 2/4/8 s. O piso '
         'da latência por salto é a cadência de 250 ms, não o intervalo '
         'de conexão.'),
    ]),
    ('Envio diferido', RGBColor(0x4F, 0x6C, 0xB0), [
        ('Fila de enlace de 32', 'Compartilhada. Um fan-out de broadcast '
         'pode ocupar um slot por enlace de uma só vez.'),
        ('Expira em 8 s', 'Cobre a escada de tentativas (100 ms << n). Só '
         'congestionamento é enfileirado: recusa estrutural falharia igual '
         'na retentativa.'),
    ]),
]
cwm = (WIDTH - 0.36) / 3
for i, (title, col, items) in enumerate(mech):
    x = LEFT + i * (cwm + 0.18)
    card(s, x, TOP + 0.24, cwm, 3.68)
    chip(s, x + 0.22, TOP + 0.48, cwm - 0.44, 0.44, title, col, size=14)
    yy = TOP + 1.14
    for head, body in items:
        textbox(s, x + 0.22, yy, cwm - 0.44, 0.30,
                [[{'t': head, 'size': 13.5, 'b': True, 'c': NAVY}]])
        textbox(s, x + 0.22, yy + 0.34, cwm - 0.44, 0.95,
                [[{'t': body, 'size': 12}]], line=1.03)
        yy += 1.36

card(s, LEFT, TOP + 4.18, WIDTH, 1.32, fill=CARD_ALT)
badge(s, LEFT + 0.26, TOP + 4.40, 0.34, '?', fill=ORANGE, size=15)
textbox(s, LEFT + 0.70, TOP + 4.43, WIDTH - 1.0, 0.3,
        [[{'t': 'Por que broadcast não é confirmado', 'size': 15, 'b': True,
           'c': NAVY}]])
textbox(s, LEFT + 0.26, TOP + 4.86, WIDTH - 0.52, 0.55, [
    [{'t': 'O remetente não tem a lista de membros da rede, então não há a '
           'quem cobrar ACK — e cada enlace físico já é confiável por '
           'GATT/CBFC. A repetição fica com o originador (s_broadcastRetry: '
           'até 3 reenvios espaçados 800 ms), não com a malha.',
      'size': 12.5}],
], line=1.05)

# ================================================================ slide 11
s = new_slide('Supervisão de enlace e parâmetros de conexão',
              'Como um enlace morto é detectado, e por que o app precisa de '
              'parâmetros próprios')
sup = [('30 s ± 5 s', 'Heartbeat', 'Token por enlace, ecoado em '
        'HEARTBEAT_ACK. Prova que as duas direções da aplicação estão vivas '
        '— o evento de conexão sozinho não detecta caminho TRSP travado.'),
       ('3 faltas', 'Desconexão', 'Três ACKs perdidos derrubam o enlace. '
        'Qualquer pacote válido também conta como presença.'),
       ('90 s', 'Enlace morto', 'MESH_LINK_DEAD_TICKS sem atividade '
        'nenhuma derruba o enlace, independentemente do heartbeat.'),
       ('30 s', 'Varredura', 'Enlace que nunca ficou utilizável é '
        'descartado — exceto o do app, que não tem JOIN a completar.')]
cws = (WIDTH - 0.54) / 4
for i, (big, head, body) in enumerate(sup):
    x = LEFT + i * (cws + 0.18)
    card(s, x, TOP + 0.24, cws, 2.30)
    textbox(s, x + 0.22, TOP + 0.44, cws - 0.44, 0.42,
            [[{'t': big, 'size': 22, 'b': True, 'c': NAVY}]])
    textbox(s, x + 0.22, TOP + 0.92, cws - 0.44, 0.28,
            [[{'t': head, 'size': 13, 'b': True, 'c': ORANGE}]])
    textbox(s, x + 0.22, TOP + 1.26, cws - 0.44, 1.05,
            [[{'t': body, 'size': 11.5}]], line=1.02)

rows_cp = [['Parâmetro', 'Enlace de malha', 'Enlace do app'],
           ['Intervalo', '40 – 80 ms', '30 – 50 ms'],
           ['Peripheral latency', '0', '4'],
           ['Supervision timeout', '20 s', '5 s'],
           ['Definido por', 'central, no CreateConnection',
            'periférico, após TX_OPENED']]
table(s, LEFT, TOP + 2.86, 7.15, 2.05, rows_cp, [2.15, 2.55, 2.45],
      first_col_bold=True, size=12)

rx = LEFT + 7.50
rw = RIGHT - rx
card(s, rx, TOP + 2.86, rw, 2.05, fill=CARD_ALT)
badge(s, rx + 0.24, TOP + 3.06, 0.32, '!', fill=ORANGE, size=14)
textbox(s, rx + 0.64, TOP + 3.08, rw - 0.88, 0.28,
        [[{'t': 'Por que o app difere', 'size': 14, 'b': True, 'c': NAVY}]])
textbox(s, rx + 0.24, TOP + 3.48, rw - 0.48, 1.30, [
    [{'t': 'Com latency 0, todo evento de conexão perdido conta contra o '
           'timeout — e o nó perde eventos: seis enlaces, anúncio '
           'permanente e janelas de scan não cabem no rádio. Latency 4 '
           'torna o salto legal em vez de fatal. E 20 s de supervision '
           'timeout está fora do que o celular aceita (a Apple limita a '
           '6 s), então o pedido era simplesmente recusado.', 'size': 11.5}],
], line=1.02)

# ================================================================ slide 12
s = new_slide('Tempos de formação de rede',
              'Orçamento derivado dos temporizadores — estimativa, não '
              'medição de bancada')
budget = [('Janela de scan', '6 s', 'MESH_SCAN_DURATION_MS'),
          ('Aquisição da conexão', '2 – 5 s', 'anúncio 200-400 ms, scan 20%'),
          ('GATT + CBFC + JOIN', '3 – 8 s', 'retry com backoff exponencial'),
          ('Uma onda, caso limpo', '11 – 19 s', 'soma dos três acima'),
          ('Rescan após falha', '+30 s', 'MESH_RESCAN_PERIOD_TICKS'),
          ('Backoff por recusa', '15 – 75 s', '15 s x nº de falhas, até 5x')]
textbox(s, LEFT, TOP + 0.16, 5.75, 0.28,
        [[{'t': 'Orçamento por onda de crescimento', 'size': 14, 'b': True,
           'c': NAVY}]])
yy = TOP + 0.54
for i, (name, val, src) in enumerate(budget):
    hl = (i == 3)
    card(s, LEFT, yy, 5.75, 0.62, fill=CARD_ALT if hl else CARD)
    textbox(s, LEFT + 0.20, yy + 0.10, 2.75, 0.26,
            [[{'t': name, 'size': 12.5, 'b': True,
               'c': ORANGE if hl else INK}]])
    textbox(s, LEFT + 0.20, yy + 0.34, 3.00, 0.22,
            [[{'t': src, 'size': 10, 'c': MUTED, 'font': MONO_FONT}]])
    textbox(s, LEFT + 3.25, yy + 0.16, 2.30, 0.32,
            [[{'_': {'align': PP_ALIGN.RIGHT}, 't': val, 'size': 15,
               'b': True, 'c': ORANGE if hl else NAVY}]])
    yy += 0.70

chart_data = CategoryChartData()
chart_data.categories = ['5 nós', '18 nós', '50 nós', '100 nós']
chart_data.add_series('Típico (s)', (15, 40, 75, 120))
chart_data.add_series('Pior caso (s)', (40, 120, 210, 300))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(LEFT + 6.05),
                        Inches(TOP + 0.30), Inches(RIGHT - LEFT - 6.05),
                        Inches(3.30), chart_data)
ch = gf.chart
ch.has_title = True
ch.chart_title.text_frame.text = 'Tempo até MESH FORMED em todos os nós'
tp = ch.chart_title.text_frame.paragraphs[0]
tp.runs[0].font.size = Pt(13)
tp.runs[0].font.bold = True
tp.runs[0].font.color.rgb = NAVY
tp.runs[0].font.name = BODY_FONT
ch.has_legend = True
ch.legend.position = XL_LEGEND_POSITION.TOP
ch.legend.include_in_layout = False
ch.legend.font.size = Pt(11)
ch.legend.font.name = BODY_FONT
ch.plots[0].series[0].format.fill.solid()
ch.plots[0].series[0].format.fill.fore_color.rgb = NAVY
ch.plots[0].series[1].format.fill.solid()
ch.plots[0].series[1].format.fill.fore_color.rgb = ORANGE
plot = ch.plots[0]
plot.has_data_labels = True
plot.data_labels.font.size = Pt(10)
plot.data_labels.font.name = BODY_FONT
plot.data_labels.font.color.rgb = INK
ch.category_axis.tick_labels.font.size = Pt(11)
ch.category_axis.tick_labels.font.name = BODY_FONT
ch.value_axis.tick_labels.font.size = Pt(10)
ch.value_axis.tick_labels.font.name = BODY_FONT
ch.value_axis.has_major_gridlines = True
ch.value_axis.major_gridlines.format.line.color.rgb = RULE
ch.value_axis.major_gridlines.format.line.width = Pt(0.75)

card(s, LEFT + 6.05, TOP + 3.76, RIGHT - LEFT - 6.05, 1.40, fill=CARD_ALT)
badge(s, LEFT + 6.29, TOP + 3.96, 0.32, 'i', fill=ORANGE, size=14)
textbox(s, LEFT + 6.69, TOP + 3.98, RIGHT - LEFT - 6.75, 0.28,
        [[{'t': 'Como ler estes números', 'size': 13.5, 'b': True,
           'c': NAVY}]])
textbox(s, LEFT + 6.29, TOP + 4.34, RIGHT - LEFT - 6.60, 0.75, [
    [{'t': 'Derivados do orçamento de temporizadores ao lado, não medidos '
           'em hardware. A formação é amplamente paralela: o gargalo é a '
           'disputa por slots dos ids baixos e a contenção de rádio, não o '
           'número de nós. Meça com a linha MESH FORMED de cada nó.',
      'size': 11.5}],
], line=1.02)

# ================================================================ slide 13
s = new_slide('Indicador de rede formada',
              'Critério local conservador, exposto em três lugares')
card(s, LEFT, TOP + 0.24, 6.30, 2.00)
badge(s, LEFT + 0.24, TOP + 0.44, 0.34, '✓', fill=GREEN, size=15)
textbox(s, LEFT + 0.68, TOP + 0.47, 5.4, 0.3,
        [[{'t': 'Critério', 'size': 15, 'b': True, 'c': NAVY}]])
textbox(s, LEFT + 0.24, TOP + 0.92, 5.85, 1.30, [
    [{'t': 'Todo enlace de malha do nó está isReady e topologyReady;'}],
    [{'t': 'existe pelo menos um enlace utilizável;'}],
    [{'t': 'o nó tem lugar na árvore (tem pai ou é a raiz);'}],
    [{'t': 'e isso se manteve por 5 s seguidos.', 'b': True}],
], default_size=12.5, line=1.05, space_after=3)

card(s, LEFT + 6.60, TOP + 0.24, RIGHT - LEFT - 6.60, 2.00, fill=CARD_ALT)
textbox(s, LEFT + 6.84, TOP + 0.44, RIGHT - LEFT - 7.10, 0.3,
        [[{'t': 'Por que o atraso de 5 s', 'size': 15, 'b': True,
           'c': NAVY}]])
textbox(s, LEFT + 6.84, TOP + 0.88, RIGHT - LEFT - 7.10, 1.30, [
    [{'t': 'Um enlace que sobe e cai em seguida produziria um par '
           'FORMED/NOT FORMED no console a cada oscilação. Exigir '
           'estabilidade contínua transforma o indicador em sinal de '
           'convergência, não em eco de cada evento de conexão.',
      'size': 12.5}],
], line=1.05)

outs = [('Console', [
    ('*** MESH FORMED *** id=3 root=1 depth=2', MONO_FONT, GREEN, True),
    ('    parent=1 links=2 nodes=5', MONO_FONT, GREEN, True),
    ('', MONO_FONT, INK, False),
    ('STATUS id=3 links=2 root=1 depth=2', MONO_FONT, INK, False),
    ('    parent=1 free=3 FORMED', MONO_FONT, INK, False)]),
    ('Anúncio BLE', [
        ('Dados de fabricante 0xFEDA, byte de flags:', BODY_FONT, INK, False),
        ('', BODY_FONT, INK, False),
        ('0x01  nó de malha', MONO_FONT, INK, False),
        ('0x02  tem lugar na árvore', MONO_FONT, INK, False),
        ('0x04  rede formada  ← novo', MONO_FONT, ORANGE, True)]),
    ('API em C', [
        ('bool MESH_IsNetworkFormed(void);', MONO_FONT, INK, False),
        ('uint8_t MESH_GetSubtreeNodeCount(void);', MONO_FONT, INK, False),
        ('', BODY_FONT, INK, False),
        ('Este nó mais tudo o que é reportado', BODY_FONT, INK, False),
        ('abaixo dele na árvore.', BODY_FONT, INK, False)])]
cwo = (WIDTH - 0.36) / 3
for i, (head, lines) in enumerate(outs):
    x = LEFT + i * (cwo + 0.18)
    card(s, x, TOP + 2.56, cwo, 2.35)
    chip(s, x + 0.22, TOP + 2.76, 1.65, 0.34, head, NAVY, size=12)
    textbox(s, x + 0.22, TOP + 3.24, cwo - 0.44, 1.50,
            [[{'t': ln, 'size': 10.5, 'font': f, 'c': c, 'b': b}]
             for ln, f, c, b in lines], line=1.10, space_after=1)

# ================================================================ slide 14
s = new_slide('Dimensionamento para 100 nós',
              'Constantes revisadas e o custo em memória estática')
rows_sz = [
    ['Constante', 'Antes', 'Depois', 'Motivo'],
    ['MESH_MAX_TTL', '12', '32', 'diâmetro do grafo, não da árvore'],
    ['MESH_DUP_WINDOW', '32', '64', 'bitmap uint64_t; 32 cobria 1,6 s'],
    ['MESH_MAX_DISCOVERED', '16', '48', 'menor que os nós ao alcance'],
    ['MESH_CHILD_SUMMARY_SIZE', '3', '4', 'o 4º filho sumia do agregado'],
    ['MESH_TX_QUEUE_SIZE', '8', '24', 'rajada de comandos unicast'],
    ['MESH_LINK_TX_QUEUE_SIZE', '16', '32',
     'fan-out de broadcast × 5 enlaces'],
    ['MESH_ACK_TIMEOUT_TICKS', '700 ms', '2000 ms', 'RTT real de 8 saltos'],
    ['Cadência MESH_Maintenance', '1000 ms', '250 ms',
     'piso da latência por salto'],
]
table(s, LEFT, TOP + 0.20, 8.45, 3.90, rows_sz, [2.60, 1.10, 1.10, 3.65],
      size=11, header_size=11, mono_cols=(0,))

rx = LEFT + 8.80
rw = RIGHT - rx
card(s, rx, TOP + 0.20, rw, 1.80, fill=CARD_ALT)
textbox(s, rx + 0.24, TOP + 0.42, rw - 0.48, 0.30,
        [[{'t': 'Custo em RAM', 'size': 15, 'b': True, 'c': NAVY}]])
textbox(s, rx + 0.24, TOP + 0.80, rw - 0.48, 0.95, [
    [{'t': '+2 459 B', 'size': 22, 'b': True, 'c': ORANGE}],
    [{'t': 'bss de 87,2 KB para 89,7 KB. Total estático de 90,4 KB '
           'em 128 KB.', 'size': 11.5}],
], line=1.02, space_after=2)

card(s, rx, TOP + 2.20, rw, 1.90)
textbox(s, rx + 0.24, TOP + 2.42, rw - 0.48, 0.30,
        [[{'t': 'REPAIR_REQUEST', 'size': 15, 'b': True, 'c': NAVY}]])
textbox(s, rx + 0.24, TOP + 2.80, rw - 0.48, 1.05, [
    [{'t': 'Deixou de ser inundação global (TTL 2). Perder um nó de nível 1 '
           'deixa ~21 nós órfãos, e 21 inundações simultâneas seriam '
           '~6 000 pacotes no pior momento possível.', 'size': 11.5}],
], line=1.02)

textbox(s, LEFT, TOP + 4.35, 8.45, 1.05, [
    [{'t': 'Não alterado: ', 'b': True, 'size': 12.5, 'c': NAVY},
     {'t': 'a fila de eventos do app segue em 64. APP_Msg_T ocupa 260 B, '
           'então ela já consome 16,6 KB dos 40 KB de heap do FreeRTOS — '
           'aumentá-la deixaria pouco para as alocações de evento do stack '
           'BLE. O gatilho real de estouro não é o volume de anúncios, é '
           'LED_Dimmer_Identify(), que bloqueia a task do app por 3 s.',
      'size': 12.5}],
], line=1.05)

prs.save(DECK)
print('deck salvo')
